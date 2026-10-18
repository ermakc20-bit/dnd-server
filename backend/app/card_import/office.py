from __future__ import annotations

import json
from io import BytesIO
from pathlib import PurePosixPath
from typing import Any
from zipfile import ZipFile

from defusedxml.ElementTree import fromstring  # type: ignore[import-untyped]
from openpyxl import load_workbook  # type: ignore[import-untyped]

from app.card_import.models import MAX_EXPANDED, SourceFragment, fail
from app.card_import.templates import SHEETS

LABELS = {
    "Внешний ключ": "external_key",
    "Название": "name",
    "Тип": "kind",
    "Количество": "count",
    "Видимость": "visibility",
    "Пакет": "system_id",
    "Подзаголовок": "subtitle",
    "Заметки": "notes",
    "Создать токен": "create_token",
}
KINDS = {"Персонаж": "player_character", "NPC": "npc", "Моб": "mob", "Объект": "object"}
VISIBILITY = {"Всем": "everyone", "Только мастеру": "gm_only", "Владельцу": "owner_only"}


def value(key: str, raw: Any) -> Any:
    if key == "kind":
        return KINDS.get(str(raw), raw)
    if key == "visibility":
        return VISIBILITY.get(str(raw), raw)
    if key == "create_token":
        if raw not in (True, False, "Да", "Нет", "true", "false"):
            raise fail("invalid_field", "Выберите Да или Нет", field=key)
        return raw in (True, "Да", "true")
    return raw


def schema_version(raw: Any, section: str, row: int) -> int:
    try:
        if isinstance(raw, bool) or (isinstance(raw, float) and not raw.is_integer()):
            raise ValueError
        version = int(raw)
    except (TypeError, ValueError, OverflowError):
        raise fail(
            "invalid_field", "Укажите целую версию", section=section, row=row,
            field="schema_version",
        ) from None
    if version not in (0, 1):
        raise fail(
            "schema_version", "Поддерживаются версии 0 и 1", section=section, row=row,
            field="schema_version",
        )
    return version


def safe_zip(data: bytes, format: str) -> dict[str, bytes]:
    result: dict[str, bytes] = {}
    with ZipFile(BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 500 or sum(e.file_size for e in entries) > MAX_EXPANDED:
            raise fail("archive_limit", "Слишком большой распакованный документ")
        for entry in entries:
            name = entry.filename
            path = PurePosixPath(name)
            if (
                path.is_absolute()
                or ".." in path.parts
                or "\\" in name
                or ":" in name
                or name in result
                or entry.flag_bits & 1
                or entry.file_size > MAX_EXPANDED
                or (entry.external_attr >> 16) & 0o170000 == 0o120000
            ):
                raise fail("unsafe_archive", "Недопустимый путь или запись архива")
            if any(
                s in name.lower()
                for s in ("vbaproject", "embeddings/", "activex/", "externallinks/")
            ):
                raise fail("active_content", "Макросы и вложения запрещены")
            if (
                not entry.is_dir()
                and not name.endswith(".rels")
                and path.suffix.lower()
                not in {
                    ".xml",
                    ".rels",
                    ".png",
                    ".jpg",
                    ".jpeg",
                }
            ):
                raise fail("active_content", "Исполняемые и неизвестные вложения запрещены")
            content = archive.read(entry)
            result[name] = content
            # OOXML relationships may target XML with any extension. Only real
            # binary image signatures are exempt; renamed XML must use the same
            # DTD/entity/active-content guards before an Office library sees it.
            image = content.startswith((b"\x89PNG\r\n\x1a\n", b"\xff\xd8\xff"))
            if not entry.is_dir() and not image:
                root = fromstring(content, forbid_dtd=True)
                for element in root.iter():
                    local = element.tag.rsplit("}", 1)[-1]
                    if (
                        element.attrib.get("TargetMode") == "External"
                        or local in {"f", "fldSimple", "instrText", "altChunk", "object"}
                        or any("macroEnabled" in text for text in element.attrib.values())
                    ):
                        raise fail(
                            "active_content",
                            "Формулы, внешние ссылки и активное содержимое запрещены",
                        )
        required = "xl/workbook.xml" if format == "xlsx" else "word/document.xml"
        if required not in result:
            raise fail("corrupt_file", "Документ не соответствует выбранному формату")
    return result


def read_xlsx(data: bytes) -> tuple[list[dict[str, Any]], int]:
    safe_zip(data, "xlsx")
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        if "Инструкция" not in workbook or "Карточки" not in workbook:
            raise fail("schema_version", "Используйте версионированный шаблон XLSX")
        if workbook["Инструкция"]["A1"].value != "schema_version":
            raise fail("schema_version", "В A1 укажите schema_version")
        version = schema_version(workbook["Инструкция"]["B1"].value, "Инструкция", 1)
        cards: dict[str, dict[str, Any]] = {}
        for sheet in sorted(workbook, key=lambda sheet: sheet.title != "Карточки"):
            if sheet.title == "Инструкция":
                continue
            if sheet.title not in SHEETS:
                raise fail(
                    "unknown_sheet",
                    "Неизвестный лист: перенесите данные в Поля",
                    section=sheet.title[:120],
                )
            if sheet.max_row > 2001 or sheet.max_column > 210:
                raise fail("row_limit", "Слишком много строк или столбцов", section=sheet.title)
            rows = sheet.iter_rows(values_only=True)
            headers = [str(cell or "") for cell in next(rows)]
            if len(set(headers)) != len(headers) or "Внешний ключ" not in headers:
                raise fail("headers", "Проверьте заголовки столбцов", section=sheet.title)
            for rownum, cells in enumerate(rows, 2):
                if not any(cell is not None for cell in cells):
                    continue
                fields = {k: v for k, v in zip(headers, cells, strict=True) if v is not None}
                external = str(fields.pop("Внешний ключ", ""))
                if not external:
                    raise fail(
                        "external_key", "Укажите внешний ключ", section=sheet.title, row=rownum
                    )
                if sheet.title == "Карточки":
                    if external in cards:
                        raise fail(
                            "duplicate_key", "Ключ повторяется", row=rownum, field="external_key"
                        )
                    cards[external] = {
                        LABELS.get(k, k): value(LABELS.get(k, k), v) for k, v in fields.items()
                    }
                    cards[external].update(external_key=external, _row=rownum)
                    cards[external]["_source"] = {
                        LABELS.get(label, label): {
                            "section": sheet.title, "row": rownum, "field": label,
                        }
                        for label in fields
                    }
                    continue
                if external not in cards:
                    raise fail(
                        "external_key",
                        "Нет карточки с этим ключом",
                        section=sheet.title,
                        row=rownum,
                        field="external_key",
                        entity=external,
                    )
                card = cards[external]
                for label, cell_value in fields.items():
                    if label not in SHEETS[sheet.title]:
                        additional = f"{sheet.title} {rownum} {label}"
                        card[additional] = str(cell_value)
                        card.setdefault("_source", {})[additional] = {
                            "section": sheet.title,
                            "row": rownum,
                            "field": label,
                        }
                visibility = VISIBILITY.get(
                    str(fields.get("Видимость", "Всем")), fields.get("Видимость", "everyone")
                )
                item: dict[str, Any] = {"name": fields.get("Название", "")}
                collection = {
                    "Ресурсы": "resources",
                    "Навыки": "skills",
                    "Инвентарь": "inventory",
                    "Способности": "abilities",
                    "Поля": "fields",
                    "Токены": "create_token",
                }[sheet.title]
                if collection == "create_token":
                    card[collection] = value(collection, fields.get("Создать токен", "Нет"))
                    continue
                if collection in {"resources", "skills", "fields"}:
                    item["key"] = fields.get("Ключ", "")
                if collection in {"resources", "fields"}:
                    item["label"] = item.pop("name")
                    item["visibility"] = visibility
                mapping = {
                    "Текущее": "current",
                    "Максимум": "max",
                    "Бонус": "bonus",
                    "Описание": "description",
                    "Количество": "quantity",
                    "Действие": "action_type",
                    "Значение": "value",
                    "Тип значения": "value_type",
                }
                item.update(
                    {target: fields[label] for label, target in mapping.items() if label in fields}
                )
                if fields.get("Геометрия JSON"):
                    item["geometry"] = json.loads(fields["Геометрия JSON"])
                card.setdefault(collection, []).append(item)
                item_path = f"{collection}.{len(card[collection]) - 1}"
                card.setdefault("_source", {})[item_path] = {
                    "section": sheet.title,
                    "row": rownum,
                    "field": collection,
                }
                for label, target in mapping.items():
                    if label in fields:
                        card["_source"][f"{item_path}.{target}"] = {
                            "section": sheet.title, "row": rownum, "field": label,
                        }
                if "geometry" in item:
                    card["_source"][f"{item_path}.geometry"] = {
                        "section": sheet.title, "row": rownum, "field": "Геометрия JSON",
                    }
        return list(cards.values()), version
    finally:
        workbook.close()


class MarkupRows(list[dict[str, Any]]):
    def __init__(self) -> None:
        super().__init__()
        self.fragments: list[SourceFragment] = []


def read_markup(text: str) -> tuple[list[dict[str, Any]], int]:
    cards = MarkupRows()
    version: int | None = None
    card: dict[str, Any] | None = None
    pending: list[str] = []
    start = 1

    def flush() -> None:
        if not pending:
            return
        if len(cards.fragments) >= 2000:
            raise fail("fragment_limit", "Слишком много нераспознанных фрагментов")
        cards.fragments.append(SourceFragment(
            section="Раздел", row=start, field="текст",
            external_key=card["external_key"] if card is not None else None,
            text="".join(pending),
        ))
        pending.clear()

    for row, line in enumerate(text.splitlines(keepends=True), 1):
        label, separator, raw = line.partition(":")
        label, raw = label.strip(), raw.strip()
        if not separator or (card is None and label not in {"schema_version", "Карточка"}):
            if not pending:
                start = row
            pending.append(line)
            continue
        flush()
        if label == "schema_version":
            version = schema_version(raw, "Раздел", row)
        elif label == "Карточка":
            if not raw:
                raise fail("external_key", "Укажите ключ карточки", section="Раздел", row=row,
                           field="Карточка")
            card = {"external_key": raw, "_section": "Раздел", "_row": row}
            cards.append(card)
        elif card is not None:
            if label in {"Ресурс", "Способность"}:
                parts = [part.strip() for part in raw.split("|")]
                if label == "Ресурс":
                    item = {
                        "key": parts[0],
                        "label": parts[1],
                        "current": parts[2],
                        "max": parts[3] if len(parts) > 3 and parts[3] else None,
                    }
                    collection = "resources"
                else:
                    item = {
                        "name": parts[0],
                        "action_type": parts[1],
                        "description": "|".join(parts[2:]),
                    }
                    collection = "abilities"
                card.setdefault(collection, []).append(item)
                field = f"{collection}.{len(card[collection]) - 1}"
            else:
                field = LABELS.get(label, label)
                if field in card:
                    raise fail("duplicate_field", "Повтор одиночного поля", row=row, field=field)
                card[field] = value(field, raw)
            card.setdefault("_source", {})[field] = {
                "section": "Раздел",
                "row": row,
                "field": label,
            }
    if version is None:
        raise fail("schema_version", "Нужна строка schema_version: 1 и заголовки Карточка: ключ")
    flush()
    return cards, version


def read_docx(data: bytes) -> tuple[list[dict[str, Any]], int]:
    parts = safe_zip(data, "docx")
    root = fromstring(parts["word/document.xml"], forbid_dtd=True)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    word = "{" + ns["w"] + "}"
    literal_controls = {word + "br": "\n", word + "cr": "\n", word + "tab": "\t"}

    def literal_text(element: Any) -> str:
        chunks: list[str] = []
        paragraph_seen = False
        stack = [(element, False)]
        while stack:
            node, in_run = stack.pop()
            if node.tag in {word + "pPr", word + "rPr"}:
                continue
            if node.tag == word + "p":
                if paragraph_seen:
                    chunks.append("\n")
                paragraph_seen = True
            if node.tag == word + "t":
                chunks.append(node.text or "")
            elif in_run and node.tag in literal_controls:
                # Run controls are literal text; formatting tab stops are not.
                chunks.append(literal_controls[node.tag])
            stack.extend((child, node.tag == word + "r") for child in reversed(list(node)))
        return "".join(chunks)

    lines = []
    body = root.find("w:body", ns)
    if body is None:
        raise fail("corrupt_file", "Пустой документ")
    for element in body:
        if element.tag.endswith("}p"):
            lines.append(literal_text(element))
        elif element.tag.endswith("}tbl"):
            for row in element.findall("w:tr", ns):
                cells = [
                    literal_text(cell)
                    for cell in row.findall("w:tc", ns)
                ]
                if len(cells) == 2:
                    lines.append(": ".join(cells))
                else:
                    lines.append("\t".join(cells))
    return read_markup("\n".join(lines))
