from io import BytesIO
from typing import Any

from docx import Document
from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]

SHEETS: dict[str, list[str]] = {
    "Карточки": ["Внешний ключ", "Название", "Тип", "Количество", "Видимость", "Пакет", "Заметки"],
    "Ресурсы": ["Внешний ключ", "Ключ", "Название", "Текущее", "Максимум", "Видимость"],
    "Навыки": ["Внешний ключ", "Ключ", "Название", "Бонус", "Описание"],
    "Инвентарь": ["Внешний ключ", "Название", "Количество", "Описание"],
    "Способности": ["Внешний ключ", "Название", "Действие", "Описание", "Геометрия JSON"],
    "Токены": ["Внешний ключ", "Создать токен"],
    "Поля": ["Внешний ключ", "Ключ", "Название", "Значение", "Тип значения", "Видимость"],
}
EXAMPLES: dict[str, list[list[Any]]] = {
    "Карточки": [
        [
            "guard",
            "Страж",
            "Моб",
            2,
            "Только мастеру",
            "neutral",
            "Пример: замените своими данными",
        ],
        ["torch", "Факел", "Объект", 1, "Всем", "neutral", ""],
    ],
    "Ресурсы": [["guard", "hp", "Здоровье", 12, 12, "Всем"]],
    "Навыки": [["guard", "notice", "Внимательность", 2, ""]],
    "Инвентарь": [["guard", "Щит", 1, ""]],
    "Способности": [["guard", "Удар щитом", "Действие", "Отталкивает цель", ""]],
    "Токены": [["guard", "Нет"]],
    "Поля": [["guard", "armor", "Защита", 13, "number", "Всем"]],
}
MARKUP = """schema_version: 1
Карточка: guard
Название: Страж
Тип: Моб
Количество: 2
Видимость: Только мастеру
Ресурс: hp | Здоровье | 12 | 12
Способность: Удар щитом | Действие | Отталкивает цель
Карточка: torch
Название: Факел
Тип: Объект
"""


def download_template(format: str) -> bytes:
    output = BytesIO()
    if format == "txt":
        return MARKUP.encode()
    if format == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Инструкция"
        for row in [
            ["schema_version", 1],
            ["Заполните примеры своими данными. Один файл — один импорт."],
            ["Внешний ключ связывает листы; он должен быть уникален в листе Карточки."],
            ["Пределы: 5 МБ, 100 карточек, 200 экземпляров. Формулы и макросы запрещены."],
            ["Видимость: Всем / Только мастеру / Владельцу. Тип: Персонаж / NPC / Моб / Объект."],
            ["Создать токен: Да / Нет. Неизвестные столбцы карточки сохраняются как поля."],
            ["Версия 0 — совместимый контракт, не исторический релиз."],
        ]:
            sheet.append(row)
        sheet.column_dimensions["A"].width = 110
        for name, headers in SHEETS.items():
            current = workbook.create_sheet(name)
            current.append(headers)
            for row in EXAMPLES[name]:
                current.append(row)
            current.freeze_panes = "A2"
            current.auto_filter.ref = current.dimensions
            for cell in current[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="344A60")
                current.column_dimensions[cell.column_letter].width = 26
        for current in workbook:
            current.page_setup.orientation = "landscape"
            current.page_setup.paperSize = current.PAPERSIZE_A4
            current.page_setup.fitToWidth = 1
            current.page_setup.fitToHeight = 0
            current.sheet_properties.pageSetUpPr.fitToPage = True
            current.print_options.horizontalCentered = True
            current.print_area = current.dimensions
            for row in current:
                current.row_dimensions[row[0].row].height = 32
                for cell in row:
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
        workbook.save(output)
    elif format == "docx":
        document = Document()
        document.add_heading("Карточки игрового стола", 0)
        document.add_paragraph("schema_version: 1")
        document.add_paragraph(
            "Замените примеры своими данными. Скопируйте раздел «Карточка» для новой сущности. "
            "Внешний ключ после заголовка должен быть уникальным. Поля «Ресурс» и «Способность» "
            "можно повторять; значения разделяются |. Формулы, вложения и ссылки запрещены."
        )
        document.add_paragraph(
            "До 5 МБ, 100 карточек и 200 экземпляров. Создание — только после проверки."
        )
        document.add_paragraph(
            "Ресурс — ключ | название | текущее | максимум. "
            "Способность — название | действие | описание. "
            "Тип — Персонаж / NPC / Моб / Объект. Количество — 1–200. "
            "Создать токен — Да / Нет. Видимость — Всем / Только мастеру / Владельцу."
        )
        table = None
        for line in MARKUP.splitlines()[1:]:
            key, value = line.split(":", 1)
            if key == "Карточка":
                document.add_heading(line, 1)
                table = document.add_table(rows=0, cols=2)
                table.style = "Light Shading Accent 1"
            elif table is not None:
                cells = table.add_row().cells
                cells[0].text, cells[1].text = key, value.strip()
        document.save(output)
    else:
        raise ValueError("Unsupported template")
    return output.getvalue()
