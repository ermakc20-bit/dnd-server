"""Explicit loopback-only test fixture. Production never imports this module."""

from __future__ import annotations

import json
import os
from collections.abc import AsyncIterator
from contextlib import asynccontextmanager
from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path
from urllib.parse import urlparse
from uuid import uuid4

import app.main as main
import asyncpg
import jwt
from app.assets.dependencies import get_asset_service
from app.assets.service import AssetService
from app.assets.storage import MemoryStorage, StorageUnavailable
from app.campaign_transfer.router import build_transfer_service
from app.campaigns.dependencies import get_campaign_service
from app.campaigns.repository import PostgresCampaignRepository
from app.campaigns.service import CampaignService
from app.card_import.templates import download_template
from app.config import Settings, get_settings
from app.scenes.dependencies import build_scene_service, get_scene_service
from app.scenes.repository import SceneRepository
from app.scenes.service import SceneService
from cryptography.hazmat.primitives.asymmetric import rsa
from fastapi import FastAPI, HTTPException, Response
from openpyxl import load_workbook
from PIL import Image

DB = os.environ["VTT_TEST_DATABASE_URL"]
REDIS = os.environ["VTT_TEST_REDIS_URL"]
for destination in (DB, REDIS):
    if urlparse(destination).hostname not in {"localhost", "127.0.0.1"}:
        raise RuntimeError("only dedicated localhost fixtures are allowed")
if urlparse(DB).path != "/vtt_ticket01":
    raise RuntimeError("explicit integration database required")
OUTPUT = Path(os.environ["VTT_E2E_OUTPUT"])
OUTPUT.mkdir(parents=True, exist_ok=True)
PORT = int(os.getenv("VTT_E2E_API_PORT", "18010"))
WEB_PORT = int(os.getenv("VTT_E2E_WEB_PORT", "5190"))
ISSUER = f"http://127.0.0.1:{PORT}/auth/v1"
KEY = rsa.generate_private_key(public_exponent=65537, key_size=2048)
PUBLIC_KEY = json.loads(jwt.algorithms.RSAAlgorithm.to_jwk(KEY.public_key()))
PUBLIC_KEY.update(kid="local-acceptance", use="sig", alg="RS256")
settings = Settings(
    environment="test",
    database_url=DB,
    redis_url=REDIS,
    supabase_url=f"http://127.0.0.1:{PORT}",
    frontend_origins=[f"http://127.0.0.1:{WEB_PORT}"],
    _env_file=None,
)
storage = MemoryStorage()
assets = AssetService(DB, storage)
main.build_asset_service = lambda *_: assets
app = main.create_app(settings)
repository = PostgresCampaignRepository(DB)
service = SceneService(SceneRepository(DB))
app.dependency_overrides[get_settings] = lambda: settings
app.dependency_overrides[get_campaign_service] = lambda: CampaignService(repository)
app.dependency_overrides[get_scene_service] = lambda: service
app.dependency_overrides[get_asset_service] = lambda: assets
fixture: dict[str, object] = {}


@app.get("/auth/v1/.well-known/jwks.json")
async def jwks() -> dict[str, object]:
    return {"keys": [PUBLIC_KEY]}


@app.get("/__e2e/bootstrap")
async def bootstrap() -> dict[str, object]:
    return {"token": fixture["token"]}


@app.get("/__test_storage/{token}")
async def blob(token: str) -> Response:
    try:
        data, mime = storage.read_grant(token)
        return Response(data, media_type=mime)
    except StorageUnavailable:
        raise HTTPException(403, "invalid private fixture grant") from None


prior = app.router.lifespan_context


@asynccontextmanager
async def lifespan(application: FastAPI) -> AsyncIterator[None]:
    connection = await asyncpg.connect(DB)
    user = uuid4()
    now = datetime.now(UTC)
    fixture["token"] = jwt.encode(
        {
            "sub": str(user),
            "email": "fixture@example.invalid",
            "iss": ISSUER,
            "aud": "authenticated",
            "iat": now,
            "exp": now + timedelta(hours=2),
        },
        KEY,
        algorithm="RS256",
        headers={"kid": "local-acceptance"},
    )
    await connection.execute(
        "insert into auth.users(id,email) values($1,$2)", user, f"acceptance-{user}@example.invalid"
    )
    image = BytesIO()
    Image.new("RGB", (512, 512), (34, 54, 65)).save(image, format="PNG")
    (OUTPUT / "ready-map.png").write_bytes(image.getvalue())
    workbook = load_workbook(BytesIO(download_template("xlsx")))
    for sheet in workbook:
        if sheet.max_row > 1:
            sheet.delete_rows(2, sheet.max_row - 1)
    workbook["Карточки"].append(["hero", "Герой теста", "Персонаж", 1, "Всем", "neutral", ""])
    workbook["Ресурсы"].append(["hero", "hp", "HP", 10, 20, "Всем"])
    workbook["Способности"].append(
        [
            "hero",
            "Сфера",
            "Действие",
            "Ручная визуализация",
            json.dumps(
                {
                    "shape": "circle",
                    "origin": "point",
                    "radius": 10,
                    "units": "фут",
                    "color": "#ff7448",
                }
            ),
        ]
    )
    workbook["Токены"].append(["hero", "Да"])
    workbook.save(OUTPUT / "filled-template.xlsx")
    async with prior(application):
        try:
            yield
        finally:
            await build_transfer_service(assets).close()
            await service.repository.engine.dispose()
            await repository._engine.dispose()
            await assets.engine.dispose()
            await build_scene_service(DB).repository.engine.dispose()
            await connection.execute(
                "delete from app.campaign_transfer_jobs where owner_user_id=$1", user
            )
            await connection.execute("delete from app.campaigns where owner_user_id=$1", user)
            await connection.execute("delete from auth.users where id=$1", user)
            left = await connection.fetchval("select count(*) from auth.users where id=$1", user)
            campaigns = await connection.fetchval(
                "select count(*) from app.campaigns where owner_user_id=$1", user
            )
            jobs = await connection.fetchval(
                "select count(*) from app.campaign_transfer_jobs where owner_user_id=$1", user
            )
            (OUTPUT / "cleanup.json").write_text(
                json.dumps(
                    {
                        "fixture_users_remaining": left,
                        "fixture_campaigns_remaining": campaigns,
                        "fixture_jobs_remaining": jobs,
                        "cleanup": "exact owned synthetic identity",
                    }
                )
            )
            await connection.close()


app.router.lifespan_context = lifespan
