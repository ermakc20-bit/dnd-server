import json
import re
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest


def test_json_is_deterministic_preserves_unknown_and_migrates_v0():
    from app.card_import.parser import parse

    raw = json.dumps(
        {
            "schema_version": 0,
            "cards": [
                {
                    "external_key": "guard",
                    "name": "Страж",
                    "kind": "mob",
                    "Знамя": "Синее",
                    "resources": [{"key": "hp", "label": "HP", "current": 12, "max": 12}],
                    "abilities": [{"name": "Удар", "description": "Явное описание"}],
                }
            ],
        }
    ).encode()
    first = parse(raw, "json")
    assert first == parse(raw, "json")
    assert first.schema_version == 1
    assert first.drafts[0].card.resources[0].current == 12
    assert first.drafts[0].card.fields[0].value == "Синее"
    warning = next(d for d in first.diagnostics if d.code == "unknown_field")
    assert (warning.entity, warning.section, warning.row, warning.field) == (
        "guard",
        "Карточки",
        1,
        "Знамя",
    )


@pytest.mark.parametrize("format", ["xlsx", "docx", "txt"])
def test_templates_roundtrip_and_explicit_markup(format):
    from app.card_import.parser import parse
    from app.card_import.templates import download_template

    result = parse(download_template(format), format)
    assert [d.card.name for d in result.drafts] == ["Страж", "Факел"]
    assert result.drafts[0].card.resources[0].current == 12
    assert result.drafts[0].card.abilities[0].name == "Удар щитом"
    assert result.drafts[0].count == 2
    assert result == parse(download_template(format), format)


def test_real_legacy_semantic_blocks_preserve_blood_and_attributes():
    from app.card_import.parser import parse

    fixtures = Path(__file__).parent / "fixtures"
    vampire = parse((fixtures / "vampires.html").read_bytes(), "html")
    old = parse((fixtures / "old-russia.html").read_bytes(), "html")
    assert vampire.drafts[0].card.resources[1].key == "blood"
    assert vampire.drafts[0].card.resources[1].current == 3
    assert vampire.drafts[0].card.inventory[0].name == "Часы"
    assert old.drafts[0].card.fields[0].value == "18"
    assert old.drafts[0].card.abilities[0].name == "Удар"
    assert vampire == parse((fixtures / "vampires.html").read_bytes(), "html")


def test_pdf_explicit_text_and_active_content_rejected():
    from app.card_import.models import ParseFailure
    from app.card_import.parser import parse
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(400, 600)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 10 500 Td (schema_version: 1) Tj ET")
    page[NameObject("/Contents")] = stream
    output = BytesIO()
    writer.write(output)
    with pytest.raises(ParseFailure) as missing:
        parse(output.getvalue(), "pdf")
    assert missing.value.diagnostics[0].code == "entity_limit"
    font[NameObject("/Subtype")] = NameObject("/Type0")
    font[NameObject("/Encoding")] = NameObject("/Identity-H")
    lines = ["schema_version: 1", "Карточка: guard", "Название: Страж", "Количество: 2"]
    stream.set_data(
        (
            "BT /F1 12 Tf 10 500 Td "
            + " ".join("<" + line.encode("utf-16-be").hex() + "> Tj 0 -20 Td" for line in lines)
            + " ET"
        ).encode()
    )
    output = BytesIO()
    writer.write(output)
    marked = parse(output.getvalue(), "pdf")
    assert marked.drafts[0].card.name == "Страж" and marked.drafts[0].count == 2
    writer.add_js("app.alert('DO_NOT_RUN')")
    output = BytesIO()
    writer.write(output)
    with pytest.raises(ParseFailure) as active:
        parse(output.getvalue(), "pdf")
    assert active.value.diagnostics[0].code == "active_content"


@pytest.mark.parametrize(
    "name,content",
    [
        ("../escaped.xml", b"<x/>"),
        ("word/vbaProject.bin", b"macro"),
        ("word/evil.exe", b"MZ"),
        ("word/extra.xml", b'<!DOCTYPE a [<!ENTITY x "boom">]><a>&x;</a>'),
        (
            "word/_rels/evil.rels",
            b'<Relationships><Relationship TargetMode="External" Target="https://invalid.test"/></Relationships>',
        ),
    ],
)
def test_office_unsafe_parts_are_rejected_before_domain_parse(name, content):
    from app.card_import.models import ParseFailure
    from app.card_import.parser import parse
    from app.card_import.templates import download_template

    output = BytesIO(download_template("docx"))
    with ZipFile(output, "a") as archive:
        archive.writestr(name, content)
    with pytest.raises(ParseFailure):
        parse(output.getvalue(), "docx")


def test_json_rejects_future_card_version_inside_current_envelope():
    from app.card_import.models import ParseFailure
    from app.card_import.parser import parse

    raw = b'{"schema_version":1,"cards":[{"schema_version":99,"name":"Future"}]}'
    with pytest.raises(ParseFailure) as error:
        parse(raw, "json")
    assert error.value.diagnostics[0].field == "schema_version"


def test_xlsx_unknown_resource_field_reports_original_sheet_and_row():
    from app.card_import.parser import parse
    from app.card_import.templates import download_template
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(download_template("xlsx")))
    workbook["Ресурсы"]["G1"] = "Цвет"
    workbook["Ресурсы"]["G2"] = "Синий"
    output = BytesIO()
    workbook.save(output)
    result = parse(output.getvalue(), "xlsx")
    warning = next(d for d in result.diagnostics if d.code == "unknown_field")
    assert (warning.entity, warning.section, warning.row, warning.field) == (
        "guard", "Ресурсы", 2, "Цвет"
    )
    assert any(field.value == "Синий" for field in result.drafts[0].card.fields)


@pytest.mark.parametrize("forbidden", ["formula", "xml_field", "dtd"])
def test_xlsx_relationship_target_is_checked_independently_of_its_extension(forbidden):
    from app.card_import.models import ParseFailure
    from app.card_import.parser import parse
    from app.card_import.templates import download_template

    def renamed_worksheet(inject):
        output = BytesIO()
        with (
            ZipFile(BytesIO(download_template("xlsx"))) as original,
            ZipFile(output, "w") as changed,
        ):
            for entry in original.infolist():
                content = original.read(entry)
                name = entry.filename
                if name in {"[Content_Types].xml", "xl/_rels/workbook.xml.rels"}:
                    content = content.replace(b"sheet2.xml", b"sheet2.png")
                if name == "xl/worksheets/sheet2.xml":
                    name = "xl/worksheets/sheet2.png"
                    if inject == "formula":
                        content = re.sub(
                            rb'<c r="B2".*?</c>', b'<c r="B2"><f>1+1</f><v>2</v></c>', content
                        )
                    elif inject == "xml_field":
                        content = content.replace(
                            b"</worksheet>", b"<instrText>DO_NOT_RUN</instrText></worksheet>"
                        )
                    elif inject == "dtd":
                        content = b"<!DOCTYPE worksheet>" + content
                changed.writestr(name, content)
        return output.getvalue()

    assert parse(renamed_worksheet(None), "xlsx").drafts[0].count == 2
    with pytest.raises(ParseFailure):
        parse(renamed_worksheet(forbidden), "xlsx")


@pytest.mark.parametrize("raw", [1.5, 1.9, "two"])
@pytest.mark.parametrize("field", ["count", "schema_version"])
@pytest.mark.parametrize("format", ["xlsx", "txt", "docx"])
def test_integer_fields_reject_fractional_or_text_values_at_their_source(raw, field, format):
    from app.card_import.models import ParseFailure
    from app.card_import.parser import parse
    from app.card_import.templates import download_template
    from docx import Document
    from openpyxl import load_workbook

    output = BytesIO()
    if format == "xlsx":
        workbook = load_workbook(BytesIO(download_template(format)))
        sheet, cell = ("Карточки", "D2") if field == "count" else ("Инструкция", "B1")
        workbook[sheet][cell] = raw
        workbook.save(output)
        expected = (
            ("guard", sheet, 2, "Количество") if field == "count" else (None, sheet, 1, field)
        )
    else:
        lines = [
            f"schema_version: {raw if field == 'schema_version' else 1}",
            "Карточка: guard",
            "Название: Страж",
            f"Количество: {raw if field == 'count' else 2}",
        ]
        if format == "docx":
            document = Document()
            for line in lines:
                document.add_paragraph(line)
            document.save(output)
        else:
            output.write("\n".join(lines).encode())
        expected = (
            ("guard", "Раздел", 4, "Количество")
            if field == "count" else (None, "Раздел", 1, field)
        )
    with pytest.raises(ParseFailure) as error:
        parse(output.getvalue(), format)
    diagnostic = error.value.diagnostics[0]
    assert diagnostic.code == "invalid_field"
    assert (diagnostic.entity, diagnostic.section, diagnostic.row, diagnostic.field) == expected


@pytest.mark.parametrize(
    "sheet,cell,field",
    [("Ресурсы", "D2", "Текущее"), ("Ресурсы", "E2", "Максимум"),
     ("Навыки", "D2", "Бонус"), ("Инвентарь", "C2", "Количество"),
     ("Поля", "D2", "Значение"), ("Способности", "E2", "Геометрия JSON")],
)
def test_xlsx_numeric_error_preserves_entity_sheet_row_and_field(sheet, cell, field):
    from app.card_import.models import ParseFailure
    from app.card_import.parser import parse
    from app.card_import.templates import download_template
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(download_template("xlsx")))
    workbook[sheet][cell] = (
        '{"shape":"circle","radius":"two"}' if field == "Геометрия JSON" else "two"
    )
    output = BytesIO()
    workbook.save(output)
    with pytest.raises(ParseFailure) as error:
        parse(output.getvalue(), "xlsx")
    diagnostic = error.value.diagnostics[0]
    assert diagnostic.code == "invalid_field"
    assert (diagnostic.entity, diagnostic.section, diagnostic.row, diagnostic.field) == (
        "guard", sheet, 2, field
    )


@pytest.mark.parametrize("format", ["txt", "docx"])
@pytest.mark.parametrize("value", ["two | 12", "12 | two"])
def test_marked_resource_numeric_error_preserves_original_location(format, value):
    from app.card_import.models import ParseFailure
    from app.card_import.parser import parse
    from docx import Document

    lines = ["schema_version: 1", "Карточка: guard", "Название: Страж",
             f"Ресурс: hp | HP | {value}"]  # noqa: RUF001 - literal Russian document field
    output = BytesIO()
    if format == "docx":
        document = Document()
        for line in lines:
            document.add_paragraph(line)
        document.save(output)
    else:
        output.write("\n".join(lines).encode())
    with pytest.raises(ParseFailure) as error:
        parse(output.getvalue(), format)
    diagnostic = error.value.diagnostics[0]
    assert diagnostic.code == "invalid_field"
    assert (diagnostic.entity, diagnostic.section, diagnostic.row, diagnostic.field) == (
        "guard", "Раздел", 4, "Ресурс"  # noqa: RUF001 - literal Russian document field
    )


def test_import_npc_defaults_are_private_and_explicit_public_is_retained():
    from app.card_import.parser import parse

    document = {
        "schema_version": 1,
        "cards": [
            {"external_key": "npc", "name": "NPC", "kind": "npc"},
            {"external_key": "mob", "name": "Mob", "kind": "mob"},
            {"external_key": "public", "name": "Public", "kind": "npc", "visibility": "everyone"},
        ],
    }
    parsed = parse(json.dumps(document).encode(), "json")
    assert [draft.card.visibility for draft in parsed.drafts] == ["gm_only", "gm_only", "everyone"]
