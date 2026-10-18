from io import BytesIO

import pytest
from app.card_import.models import ParseFailure
from app.card_import.parser import parse
from docx import Document

MARKED = (
    "  До карточек: исходный контекст  \r\n"
    "schema_version: 1\r\n"
    "Карточка: guard\r\n"
    "Название: Страж\r\n"
    "Количество: 2\r\n"
    "Ресурс: hp | HP | 12 | 12\r\n"  # noqa: RUF001 - literal Russian markup
    "Знамя: Синее\r\n"
    "  Тайная история без разметки  \r\n\r\n"
    "Карточка: torch\r\n"
    "Название: Факел\r\n"
    "Тип: Объект\r\n"
    "Последний неразмеченный абзац"
)


def test_markup_keeps_exact_unrecognized_text_and_association_outside_cards():
    result = parse(MARKED.encode(), "txt")
    assert result.model_dump()["fragments"] == [
        {"section": "Раздел", "row": 1, "field": "текст", "external_key": None,
         "text": "  До карточек: исходный контекст  \r\n"},
        {"section": "Раздел", "row": 8, "field": "текст", "external_key": "guard",
         "text": "  Тайная история без разметки  \r\n\r\n"},
        {"section": "Раздел", "row": 13, "field": "текст", "external_key": "torch",
         "text": "Последний неразмеченный абзац"},
    ]
    assert result == parse(MARKED.encode(), "txt")
    assert len(result.drafts) == 2 and result.drafts[0].count == 2
    assert result.drafts[0].card.resources[0].current == 12
    assert result.drafts[0].card.fields[0].value == "Синее"
    warning = next(d for d in result.diagnostics if d.code == "unknown_field")
    assert (warning.entity, warning.row, warning.field) == ("guard", 7, "Знамя")
    assert "Тайная история" not in result.drafts[0].card.model_dump_json()


def test_docx_uses_same_source_fragment_pipeline():
    document = Document()
    for line in ["schema_version: 1", "Карточка: a", "Название: Страж",
                 "  Неизвестный абзац  "]:
        document.add_paragraph(line)
    output = BytesIO()
    document.save(output)
    result = parse(output.getvalue(), "docx")
    assert result.model_dump()["fragments"] == [
        {"section": "Раздел", "row": 4, "field": "текст", "external_key": "a",
         "text": "  Неизвестный абзац  "}
    ]


def test_fragment_metadata_does_not_swallow_an_explicit_unknown_field():
    result = parse(("schema_version: 1\nКарточка: a\nНазвание: Страж\n"  # noqa: RUF001 - literal Russian markup
                    "_fragments: Явное дополнительное поле\nНеразмеченный абзац").encode(), "txt")  # noqa: RUF001 - literal Russian markup
    assert result.drafts[0].card.fields[0].label == "_fragments"
    assert result.drafts[0].card.fields[0].value == "Явное дополнительное поле"
    assert result.model_dump()["fragments"][0]["text"] == "Неразмеченный абзац"


def test_docx_unrecognized_table_rows_are_not_silently_lost():
    document = Document()
    for line in ["schema_version: 1", "Карточка: a", "Название: Страж"]:
        document.add_paragraph(line)
    document.add_table(rows=1, cols=1).cell(0, 0).text = "  Неизвестная строка таблицы  "
    output = BytesIO()
    document.save(output)
    result = parse(output.getvalue(), "docx")
    assert result.model_dump()["fragments"][0]["text"] == "  Неизвестная строка таблицы  "


@pytest.mark.parametrize("container", ["paragraph", "table_cell"])
def test_docx_literal_breaks_and_tabs_survive_in_source_fragments(container):
    document = Document()
    for line in ["schema_version: 1", "Карточка: a", "Название: Страж"]:
        document.add_paragraph(line)
    paragraph = (document.add_paragraph() if container == "paragraph"
                 else document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0])
    run = paragraph.add_run("  before")
    run.add_tab()
    run.add_text("after tab")
    run.add_break()
    run.add_text("next")
    run.add_tab()
    run.add_text("  end  ")
    output = BytesIO()
    document.save(output)
    result = parse(output.getvalue(), "docx")
    assert result.model_dump()["fragments"] == [
        {"section": "Раздел", "row": 4, "field": "текст", "external_key": "a",
         "text": "  before\tafter tab\nnext\t  end  "}
    ]
    assert len(result.drafts) == 1 and result.drafts[0].card.name == "Страж"
    assert result.drafts[0].card.notes is None


def test_docx_cell_paragraph_boundaries_are_preserved_with_literal_breaks():
    document = Document()
    for line in ["schema_version: 1", "Карточка: a", "Название: Страж"]:
        document.add_paragraph(line)
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].add_run(" first ")
    run = cell.add_paragraph().add_run(" second")
    run.add_tab()
    run.add_text("column")
    run.add_break()
    run.add_text(" third ")
    output = BytesIO()
    document.save(output)
    result = parse(output.getvalue(), "docx")
    assert result.model_dump()["fragments"][0]["text"] == " first \n second\tcolumn\n third "


def test_long_unrecognized_text_is_kept_without_card_text_limit_truncation():
    literal = "я" * 10001
    result = parse(("schema_version: 1\nКарточка: a\nНазвание: Страж\n" + literal).encode(), "txt")  # noqa: RUF001 - literal Russian markup
    assert result.model_dump()["fragments"][0]["text"] == literal
    assert result.drafts[0].card.notes is None


def test_blank_marked_entity_key_cannot_leave_source_fragments_unassociated():
    with pytest.raises(ParseFailure) as failure:
        parse("schema_version: 1\nКарточка: \nНазвание: Страж\nНеразмеченный абзац".encode(), "txt")  # noqa: RUF001 - literal Russian markup
    assert failure.value.diagnostics[0].code == "external_key"
    assert failure.value.diagnostics[0].row == 2


@pytest.mark.parametrize("source", ["История о стражах и факелах",  # noqa: RUF001 - literal Russian markup
                                    "schema_version: 1\nИстория о стражах и факелах"])  # noqa: RUF001 - literal Russian markup
def test_unmarked_prose_never_invents_entities(source):
    with pytest.raises(ParseFailure) as failure:
        parse(source.encode(), "txt")
    assert failure.value.diagnostics[0].code in {"schema_version", "entity_limit"}
