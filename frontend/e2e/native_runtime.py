"""Disposable local acceptance API; never connect this fixture to hosted databases."""
from contextlib import asynccontextmanager
from datetime import UTC, datetime, timedelta
from io import BytesIO
import json
import os
from pathlib import Path
from uuid import uuid4

import asyncpg
from fastapi import Response
from openpyxl import load_workbook
from PIL import Image

import app.auth.dependencies as auth_dependencies
import app.campaigns.router as campaign_router
import app.identity.dependencies as identity_dependencies
import app.main as main
from app.assets.dependencies import get_asset_service
from app.assets.service import AssetService
from app.assets.storage import MemoryStorage
from app.auth.jwt import AuthenticationError, AuthIdentity
from app.card_import.templates import download_template
from app.campaigns.dependencies import get_campaign_service
from app.campaigns.repository import PostgresCampaignRepository
from app.campaigns.service import CampaignService
from app.config import Settings, get_settings
from app.scenes.dependencies import get_scene_service
from app.scenes.repository import SceneRepository
from app.scenes.service import SceneService

DB = os.environ['VTT_TEST_DATABASE_URL']
if not DB.startswith('postgresql://postgres@127.0.0.1:55439/'):
    raise RuntimeError('Acceptance fixture requires the dedicated local test database')
OUTPUT = Path(os.environ.get('VTT_E2E_OUTPUT', '/private/tmp/dnd-t09-e2e'))
OUTPUT.mkdir(parents=True, exist_ok=True)
if not os.environ['VTT_TEST_REDIS_URL'].startswith('redis://127.0.0.1:56439/'):
    raise RuntimeError('Acceptance fixture requires dedicated local test Redis')
settings = Settings(environment='test', database_url=DB,
                    redis_url=os.environ['VTT_TEST_REDIS_URL'],
                    frontend_origins=['http://127.0.0.1:5189'], _env_file=None)
storage = MemoryStorage()
assets = AssetService(DB, storage)
main.build_asset_service = lambda *_: assets
app = main.create_app(settings)
repository = PostgresCampaignRepository(DB)
service = SceneService(SceneRepository(DB))
app.dependency_overrides[get_settings] = lambda: settings
app.dependency_overrides[get_campaign_service] = lambda: CampaignService(repository)
app.dependency_overrides[get_scene_service] = lambda: service
app.dependency_overrides[get_asset_service] = lambda: assets
fixture = {}

class FixtureVerifier:
    def verify(self, token):
        if token != 'local-t09-gm':
            raise AuthenticationError('Unknown local fixture identity')
        return AuthIdentity(user_id=fixture['user'], email='fixture@example.invalid',
                            audience='authenticated', expires_at=datetime.now(UTC) + timedelta(hours=2))

# Only JWT verification is replaced. Cookie join/restore, policies, commands and WS are real.
auth_dependencies.get_verifier = lambda _: FixtureVerifier()
campaign_router.get_verifier = lambda _: FixtureVerifier()
identity_dependencies.get_verifier = lambda _: FixtureVerifier()

async def sign(key, _ttl):
    return '/__e2e/storage/' + key
storage.sign_download = sign

@app.get('/__e2e/storage/{key:path}')
async def blob(key: str):
    data, mime = storage.objects[key]
    return Response(data, media_type=mime)

prior = app.router.lifespan_context
@asynccontextmanager
async def lifespan(application):
    connection = await asyncpg.connect(DB)
    user = uuid4()
    fixture['user'] = user
    await connection.execute('insert into auth.users(id,email) values($1,$2)', user, f't09-e2e-{user}@example.invalid')
    image = BytesIO()
    Image.new('RGB', (512, 512), (34, 54, 65)).save(image, format='PNG')
    (OUTPUT / 'ready-map.png').write_bytes(image.getvalue())
    workbook = load_workbook(BytesIO(download_template('xlsx')))
    for name in ['Карточки', 'Ресурсы', 'Навыки', 'Инвентарь', 'Способности', 'Токены', 'Поля']:
        sheet = workbook[name]
        if sheet.max_row > 1: sheet.delete_rows(2, sheet.max_row - 1)
    workbook['Карточки'].append(['hero', 'Герой теста', 'Персонаж', 1, 'Всем', 'neutral', ''])
    workbook['Ресурсы'].append(['hero', 'hp', 'HP', 10, 20, 'Всем'])
    workbook['Способности'].append(['hero', 'Сфера', 'Действие', 'Ручная визуализация', json.dumps({'shape': 'circle', 'origin': 'point', 'radius': 10, 'units': 'фут', 'color': '#ff7448'})])
    workbook['Токены'].append(['hero', 'Да'])
    workbook.save(OUTPUT / 'filled-template.xlsx')
    async with prior(application):
        try:
            yield
        finally:
            await service.repository.engine.dispose()
            await repository._engine.dispose()
            await assets.engine.dispose()
            await connection.execute('delete from app.campaign_transfer_jobs where owner_user_id=$1', user)
            await connection.execute('delete from app.campaigns where owner_user_id=$1', user)
            await connection.execute('delete from auth.users where id=$1', user)
            left = await connection.fetchval('select count(*) from auth.users where id=$1', user)
            campaigns = await connection.fetchval('select count(*) from app.campaigns where owner_user_id=$1', user)
            jobs = await connection.fetchval('select count(*) from app.campaign_transfer_jobs where owner_user_id=$1', user)
            (OUTPUT / 'cleanup.json').write_text(json.dumps({'fixture_users_remaining': left, 'fixture_campaigns_remaining': campaigns, 'fixture_jobs_remaining': jobs, 'cleanup': 'exact synthetic owner'}))
            await connection.close()
app.router.lifespan_context = lifespan
